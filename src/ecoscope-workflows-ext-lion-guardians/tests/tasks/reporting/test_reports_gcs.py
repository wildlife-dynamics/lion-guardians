"""The report tasks read their inputs from, and save their output to, GCS (ecoscope-web)."""

import io

import docx
import pandas as pd
from PIL import Image
import pytest
from cloudpathlib.local import LocalGSClient, LocalGSPath
from ecoscope_workflows_ext_custom.storage import read_bytes, write_bytes
from ecoscope_workflows_ext_custom.tasks.io import _storage

from ecoscope_workflows_ext_lion_guardians.tasks.reporting import create_vehicles_context, render_docx_page

ROOT = "gs://bucket/user/run"


@pytest.fixture
def fake_gcs(tmp_path, monkeypatch):
    """Back gs:// paths with a local directory instead of real GCS."""
    client = LocalGSClient(local_storage_dir=tmp_path / "gcs")
    monkeypatch.setattr(_storage, "_gs_path", lambda uri: LocalGSPath(uri, client=client))
    return client


def _template(text: str) -> bytes:
    doc = docx.Document()
    doc.add_paragraph(text)
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


def _text(path: str) -> str:
    return "\n".join(p.text for p in docx.Document(io.BytesIO(read_bytes(path))).paragraphs)


def test_render_docx_page_on_gcs(fake_gcs):
    template = write_bytes(_template("Name: {{ name }}"), ROOT, "template.docx")

    out = render_docx_page(template_path=template, output_dir=ROOT, context={"name": "Lion"}, filename="page.docx")

    assert out.endswith("page.docx")
    assert _text(out) == "Name: Lion"


def test_render_docx_page_embeds_gcs_image(fake_gcs):
    buf = io.BytesIO()
    Image.new("RGB", (40, 20), "red").save(buf, format="PNG")
    image = write_bytes(buf.getvalue(), ROOT, "map.png")
    template = write_bytes(_template("{{ map }}"), ROOT, "template.docx")

    out = render_docx_page(
        template_path=template, output_dir=ROOT, context={"map": image}, filename="page.docx", strict_images=True
    )

    doc = docx.Document(io.BytesIO(read_bytes(out)))
    assert len(doc.element.body.xpath(".//pic:pic")) == 1


def test_render_docx_page_escapes_values(tmp_path):
    template = tmp_path / "template.docx"
    template.write_bytes(_template("Name: {{ name }}"))

    out = render_docx_page(
        template_path=str(template), output_dir=str(tmp_path), context={"name": "<b>Lion</b> & Co"}, filename="p.docx"
    )

    assert _text(out) == "Name: <b>Lion</b> & Co"


def test_create_vehicles_context_reads_csv_on_gcs(fake_gcs):
    csv = (
        pd.DataFrame(
            {
                "subject_name": ["KDA 123"],
                "min_speed": [1.0],
                "mean_speed": [2.5],
                "max_speed": [4.0],
                "total_distance": [10.0],
            }
        )
        .to_csv(index=False)
        .encode()
    )
    summary = write_bytes(csv, ROOT, "summary.csv")

    ctx = create_vehicles_context(summary_table=summary, speedmap=None, tracks_map=None, line_chart=None)

    assert "KDA 123" in ctx.values()
